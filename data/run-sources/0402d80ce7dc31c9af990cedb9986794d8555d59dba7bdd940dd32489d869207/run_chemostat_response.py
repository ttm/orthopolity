#!/usr/bin/env python3
"""Freeze and evaluate finite-group chemostat resource-response forecasts.

Stages run in order:
  freeze    decode the calibration tables, development vessels and held-out
            baselines, then retain every forecast. No held-out post-pulse
            outcome cell is decoded.
  evaluate  verify and replay the freeze, decode held-out outcomes, score them.
  audit     replay freeze and evaluation offline against retained outputs.

Existing public measurements only; no new observations are collected.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import platform
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from orthopolity.chemostat_response import (  # noqa: E402
    baseline_anchored_forecast, common_comparison_support, cost_ratio_forecast,
    fit_cost_ratio_strength, fractional_trajectory, interpolate_trajectory,
    log_share_growth, mean_unit_trajectory, profile_tv_envelope, recovery_endpoint,
    resource_profile, time_weighted_mean, trajectory_scores, whole_unit_partition,
)
from orthopolity.gated_xlsx import numeric, read_gated_rows, serial_date  # noqa: E402

CONFIG = ROOT / 'configs/chemostat_response_2026-10-02.json'
AMENDMENTS = [ROOT / 'configs/chemostat_response_2026-10-02_amendment-1.json',
              ROOT / 'configs/chemostat_response_2026-10-02_amendment-2.json']
SOURCES = ROOT / 'data/chemostat-sources/2026-10-02'
WORKBOOKS = SOURCES / 'zenodo/extracted/scripts_data_plankton_responses_Npulse/data'
DATA = ROOT / 'data/chemostat-response/2026-10-02'
RESULTS = ROOT / 'results/chemostat-response'
SOURCE_PATHS = [Path('experiments/run_chemostat_response.py'),
                Path('src/orthopolity/chemostat_response.py'),
                Path('src/orthopolity/gated_xlsx.py')]
INPUTS = dict(main='Chemostat_experimental_timeseries.xlsx',
              stoichiometry='algae_stoichiometry_preliminary_experiments.xlsx',
              no_herbivore='algae_no-rotifer_chemostat_preliminary_experiments.xlsx')

VOLUME = 'µm³'
META = ['Category', 'Treatment', 'Chemostat', 'Date', 'Time standardised to pulse']
TOTAL_HEADER = f'Biovolume_total_algae {VOLUME}/mL'
STOICHIOMETRY_META = ['Date', 'Time standardised to pulse [days]', 'Algae']
STOICHIOMETRY_VALUES = ['N [pmol/cell]', 'C [pmol/cell]', 'C/N Ratio [molar]',
                        f'Cell volume [{VOLUME}]', 'Density [cells/mL]', f'Biovolume [{VOLUME}/mL]']
# The authors' README: Monoraphidium (Mo) and Chlorella (Co) are summed in the
# main series; "Mi" is the four-species no-herbivore mixed culture.
GROUP_SPECIES = {'Cryptomonas': 'Cr', 'Chlamydomonas': 'Ca'}
POOLED_SPECIES = ('Mo', 'Co')
MIXTURE_TREATMENT = 'Mi'
# Published main protocol (docs/chemostat-source-audit.md): N-reduced inflow.
INFLOW_NITROGEN_UMOL_PER_L = 80.
SECONDARY_MODEL = 'two_budget_cost_ratio'
ANCHORED_MODELS = ['development_mean_response', 'preliminary_no_herbivore']
TOTAL_MODELS = ['persistence', 'development_mean_response', 'preliminary_no_herbivore']
STRENGTHS = np.concatenate([[0.], np.geomspace(1e-4, 1e4, 801)])


def now():
    return datetime.now(timezone.utc).isoformat()


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def ref(path):
    path = Path(path)
    return dict(path=path.relative_to(ROOT).as_posix(), sha256=digest(path))


def read_json(path):
    return json.loads(Path(path).read_text())


def save_json(path, value):
    path = Path(path)
    content = (json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False, allow_nan=False) + '\n').encode()
    if path.exists():
        if path.read_bytes() != content:
            raise RuntimeError(f'Refusing to overwrite changed retained artifact: {path}')
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)


def verify_reference(reference):
    if digest(ROOT / reference['path']) != reference['sha256']:
        raise RuntimeError(f'Frozen reference changed: {reference["path"]}')


def jsonable(value):
    """JSON-exact structure: NaN becomes null and an extinction -inf a string."""
    if isinstance(value, dict):
        return {str(key): jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [jsonable(item) for item in value]
    if isinstance(value, np.ndarray):
        return jsonable(value.tolist())
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        value = float(value)
        if np.isnan(value):
            return None
        if np.isinf(value):
            return '-inf' if value < 0 else '+inf'
        return value
    if value is None or isinstance(value, str):
        return value
    raise TypeError(f'Unsupported value for retained JSON: {type(value).__name__}')


def as_array(values):
    return np.array([[np.nan if item is None else item for item in row] if isinstance(row, list)
                     else (np.nan if row is None else row) for row in values], dtype=float)


def composition_models(config):
    return list(config['models']) + [SECONDARY_MODEL]


def group_headers(config):
    return [f'Biovolume_{group} {VOLUME}/mL' for group in config['groups']]


def read_tables(workbooks, config, view):
    """Decode workbooks; the freeze view gates held-out post-pulse outcomes."""
    if view not in {'freeze', 'evaluation'}:
        raise ValueError('Unknown workbook view')
    development, evaluation = config['development_category'], config['evaluation_category']

    def gate(metadata):
        category, time = metadata['Category'], metadata['Time standardised to pulse']
        if category == development:
            return True
        if category == evaluation:
            if isinstance(time, bool) or not isinstance(time, (int, float)):
                raise ValueError('Held-out row without a numeric standardized time')
            return view == 'evaluation' or time <= 0
        raise ValueError(f'Undeclared vessel category: {category!r}')

    values = [TOTAL_HEADER] + group_headers(config)
    return dict(main=read_gated_rows(workbooks['main'], META, values, gate),
                stoichiometry=read_gated_rows(workbooks['stoichiometry'], STOICHIOMETRY_META,
                                              STOICHIOMETRY_VALUES, lambda metadata: True),
                no_herbivore=read_gated_rows(workbooks['no_herbivore'], META, values,
                                             lambda metadata: True))


def chemostat_id(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        return str(int(value))
    if isinstance(value, str) and value.strip():
        return value.strip()
    raise ValueError(f'Invalid chemostat identifier: {value!r}')


def parse_rows(table, config):
    records = []
    for row in table['rows']:
        metadata = row['metadata']
        time = metadata['Time standardised to pulse']
        if isinstance(time, bool) or not isinstance(time, (int, float)) or not np.isfinite(time):
            raise ValueError(f'Invalid standardized time in source row {row["source_row"]}')
        record = dict(source_row=row['source_row'], category=metadata['Category'],
                      treatment=metadata['Treatment'], chemostat=chemostat_id(metadata['Chemostat']),
                      time=float(time), decoded=row['decoded'])
        if row['decoded']:
            context = f'source row {row["source_row"]}'
            record['biovolumes'] = [numeric(row['values'][header], context) for header in group_headers(config)]
            record['total_algae'] = numeric(row['values'][TOTAL_HEADER], context)
        records.append(record)
    return records


def vessel_map(records):
    vessels = {}
    for record in records:
        vessel = vessels.setdefault(record['chemostat'], dict(
            id=record['chemostat'], category=record['category'], treatment=record['treatment'], rows=[]))
        if (vessel['category'], vessel['treatment']) != (record['category'], record['treatment']):
            raise ValueError(f'Vessel {record["chemostat"]} has inconsistent design metadata')
        vessel['rows'].append(record)
    for vessel in vessels.values():
        vessel['rows'].sort(key=lambda row: row['time'])
        times = [row['time'] for row in vessel['rows']]
        if len(times) != len(set(times)):
            raise ValueError(f'Duplicate standardized times in vessel {vessel["id"]}')
    return vessels


def complete_volumes(row):
    volumes = row.get('biovolumes')
    return volumes is not None and all(value is not None and value >= 0 for value in volumes)


def baseline_index(rows, rho):
    """Latest complete positive-total observation at or before the pulse."""
    found = None
    for index, row in enumerate(rows):
        if row['time'] <= 0 and complete_volumes(row) and float(np.dot(rho, row['biovolumes'])) > 0:
            found = index
    return found


def calibrate(table, config):
    """Pre-pulse resource per cell volume; post-pulse values are diagnostics only."""
    limit = config['calibration_time_max_days']
    species = config['calibration_species']
    densities = {'nitrogen': {}, 'carbon': {}}
    ratio_nitrogen, used, postpulse = {}, defaultdict(list), defaultdict(list)
    rows = []
    for row in table['rows']:
        values = row['values']
        context = f'stoichiometry row {row["source_row"]}'
        rows.append(dict(source_row=row['source_row'], species=row['metadata']['Algae'],
                         time=numeric(row['metadata']['Time standardised to pulse [days]'], context),
                         date=serial_date(row['metadata']['Date'], table['date1904']),
                         nitrogen=numeric(values['N [pmol/cell]'], context),
                         carbon=numeric(values['C [pmol/cell]'], context),
                         carbon_nitrogen=numeric(values['C/N Ratio [molar]'], context),
                         cell_volume=numeric(values[f'Cell volume [{VOLUME}]'], context)))
    unknown = sorted({row['species'] for row in rows} - set(species))
    if unknown:
        raise ValueError(f'Undeclared calibration species: {unknown}')
    for name in species:
        candidates = [row for row in rows if row['species'] == name]
        early = [row for row in candidates if row['time'] is not None and row['time'] <= limit]
        for resource in ('nitrogen', 'carbon'):
            values = [row[resource]/row['cell_volume'] for row in early
                      if row[resource] is not None and row['cell_volume'] is not None
                      and row[resource] > 0 and row['cell_volume'] > 0]
            if not values:
                raise ValueError(f'No eligible pre-pulse {resource} quota for {name}')
            densities[resource][name] = float(np.median(values))
        derived = [(row['carbon']/row['carbon_nitrogen'])/row['cell_volume'] for row in early
                   if row['carbon'] and row['carbon_nitrogen'] and row['cell_volume']
                   and row['carbon'] > 0 and row['carbon_nitrogen'] > 0 and row['cell_volume'] > 0]
        ratio_nitrogen[name] = float(np.median(derived)) if derived else None
        used[name] = [row['source_row'] for row in early]
        postpulse[name] = [dict(time=row['time'],
                                nitrogen_density=(row['nitrogen']/row['cell_volume'] if row['nitrogen'] and row['cell_volume'] else None),
                                carbon_density=(row['carbon']/row['cell_volume'] if row['carbon'] and row['cell_volume'] else None))
                           for row in candidates if row['time'] is not None and row['time'] > limit]
    ratios = {name: densities['carbon'][name]/densities['nitrogen'][name] for name in species}
    return dict(units='pmol per cubic micrometre of cell volume', densities=densities,
                cost_ratios_carbon_per_nitrogen=ratios, rows=rows, rows_used=dict(used),
                diagnostics=dict(
                    ratio_derived_nitrogen_densities=ratio_nitrogen,
                    ratio_derived_scope='N = C/(C:N ratio); a quota-resolution diagnostic only, never a frozen estimator',
                    postpulse_densities=dict(postpulse),
                    postpulse_scope='Preliminary-culture densities after the pulse; not used by any forecast'))


def variants(calibration, config):
    densities = calibration['densities']
    ratio = calibration['cost_ratios_carbon_per_nitrogen']
    middle_ratio = (sum(densities['carbon'][name] for name in POOLED_SPECIES)
                    / sum(densities['nitrogen'][name] for name in POOLED_SPECIES))
    pooled_index = config['groups'].index('Monoraphidium_Chlorella')
    result = []
    for resource in config['resources']:
        density = densities[resource]
        low, high = sorted(POOLED_SPECIES, key=lambda name: (density[name], name))
        for scenario in config['pooled_density_scenarios']:
            pooled = {'lower': density[low], 'upper': density[high],
                      'midpoint': .5*(density[low]+density[high])}[scenario]
            pooled_ratio = {'lower': ratio[low], 'upper': ratio[high], 'midpoint': middle_ratio}[scenario]
            vector = [density[GROUP_SPECIES[group]] for group in config['groups'] if group in GROUP_SPECIES]
            vector.insert(pooled_index, pooled)
            ratios = [ratio[GROUP_SPECIES[group]] for group in config['groups'] if group in GROUP_SPECIES]
            ratios.insert(pooled_index, pooled_ratio)
            result.append(dict(variant_id=f'{resource}-{scenario}', resource=resource, scenario=scenario,
                               densities=vector, pooled_bounds=[density[low], density[high]],
                               pooled_bound_species=[low, high], cost_ratios=ratios))
    ratios = [ratio[GROUP_SPECIES[group]] for group in config['groups'] if group in GROUP_SPECIES]
    ratios.insert(pooled_index, middle_ratio)
    result.append(dict(variant_id='biovolume-reference', resource='biovolume', scenario='unit',
                       densities=[1.]*len(config['groups']), pooled_bounds=None,
                       pooled_bound_species=None, cost_ratios=ratios))
    return result


def unit_trajectory(vessel, rho):
    index = baseline_index(vessel['rows'], rho)
    if index is None:
        return None
    rows = vessel['rows'][index:]
    if not all(row['decoded'] for row in rows):
        raise RuntimeError(f'Trajectory requested from an undecoded vessel: {vessel["id"]}')
    times = np.array([row['time'] for row in rows])
    volumes = np.array([[np.nan if value is None else value for value in row['biovolumes']] for row in rows])
    result = fractional_trajectory(times, volumes, rho, rows[0]['biovolumes'])
    values = np.full((len(times), len(rho)+1), np.nan)
    valid = result['valid']
    values[valid, :-1] = result['share_change'][valid]
    values[valid, -1] = np.log(result['total_factor'][valid])
    return dict(unit_id=vessel['id'], times=times, values=values, shares=resource_profile(volumes, rho)['shares'],
                baseline_time=float(times[0]), baseline_source_row=rows[0]['source_row'],
                baseline_shares=result['baseline_shares'], baseline_total=result['baseline_total'])


def mean_response(units, grid, minimum):
    result = mean_unit_trajectory(units, grid)
    mean = result['mean'].copy()
    available = result['contributors'] >= minimum
    mean[~available] = np.nan
    return dict(share_change=mean[:, :-1], log_total_factor=mean[:, -1],
                contributors=result['contributors'], available=available,
                unit_ids=result['unit_ids'], minimum_contributors=minimum)


def fit_strengths(units, grid, ratios, minimum):
    baselines = np.array([unit['baseline_shares'] for unit in units])
    shares = np.array([interpolate_trajectory(unit['times'], unit['shares'], grid) for unit in units])
    strength, persistence, fitted, contributors = [], [], [], []
    for index in range(len(grid)):
        result = fit_cost_ratio_strength(baselines, shares[:, index], ratios, STRENGTHS)
        enough = result['units'] >= minimum
        strength.append(result['strength'] if enough else None)
        fitted.append(result['mean_total_variation'] if enough else None)
        persistence.append(result['persistence_total_variation'] if enough else None)
        contributors.append(result['units'])
    return dict(strength=strength, development_mean_total_variation=fitted,
                development_persistence_total_variation=persistence, contributors=contributors,
                candidates='0 and 801 log-spaced values on [1e-4, 1e4]; first minimum')


def heldout_forecast(vessel, rho, ratios, development, mixture, strengths, grid, domain):
    index = baseline_index(vessel['rows'], rho)
    if index is None:
        raise ValueError(f'Held-out vessel {vessel["id"]} has no eligible baseline')
    base = vessel['rows'][index]
    initial = resource_profile(base['biovolumes'], rho)
    shares, total = initial['shares'], float(initial['total'])
    times = [row['time'] for row in vessel['rows'] if domain[0] <= row['time'] <= domain[1]]
    positions = []
    for time in times:
        if time not in grid:
            raise ValueError(f'Observation day {time} is not on the prediction grid')
        positions.append(grid.index(time))
    count, groups = len(times), len(rho)
    models = {'persistence': dict(shares=np.tile(shares, (count, 1)), total=np.full(count, total)),
              'equal_group_stock': dict(shares=np.full((count, groups), 1./groups), total=None)}
    for name, response in (('development_mean_response', development), ('preliminary_no_herbivore', mixture)):
        forecast = np.full((count, groups), np.nan)
        stock = np.full(count, np.nan)
        distance = np.full(count, np.nan)
        for row, position in enumerate(positions):
            if response['available'][position]:
                anchored = baseline_anchored_forecast(shares, response['share_change'][position])
                forecast[row], distance[row] = anchored['shares'], anchored['projection_distance']
                stock[row] = total*np.exp(response['log_total_factor'][position])
        models[name] = dict(shares=forecast, total=stock, projection_distance=distance)
    forecast = np.full((count, groups), np.nan)
    for row, position in enumerate(positions):
        if strengths['strength'][position] is not None:
            forecast[row] = cost_ratio_forecast(shares, ratios, strengths['strength'][position])
    models[SECONDARY_MODEL] = dict(shares=forecast, total=None)
    return dict(times=times, baseline_time=base['time'], baseline_source_row=base['source_row'],
                baseline_shares=shares, baseline_total=total, models=models,
                two_budget_capacity_total_variation=two_budget_capacity(shares, ratios))


def two_budget_capacity(shares, ratios):
    """Largest TV from baseline the two-budget rule reaches over its candidates."""
    weighted = np.asarray(shares)[None, :]/(1.+STRENGTHS[:, None]*np.asarray(ratios)[None, :])
    forecast = weighted/np.sum(weighted, axis=1, keepdims=True)
    return float(np.max(.5*np.sum(np.abs(forecast-np.asarray(shares)[None, :]), axis=1)))


def observed_endpoints(times, shares, baseline, ratios, config):
    """Recovery status and cost-ratio ordinal check for one observed trajectory."""
    margin = config['recovery_margin_tv']
    result = recovery_endpoint(times, shares, baseline, margin, 0., config['recovery_consecutive_observations'])
    recovery = dict(status=result['status'], confirmation_time=result['confirmation_time'],
                    within_margin_from=result.get('within_margin_from'))
    complete = np.all(np.isfinite(shares), axis=1)
    distance = .5*np.sum(np.abs(shares[complete]-baseline), axis=1)
    departed = bool(np.any(distance > margin))
    largest = float(np.max(distance)) if distance.size else None
    growth = None
    if np.all(np.asarray(baseline) > 0):
        rows = log_share_growth(shares, baseline)
        growth = [time_weighted_mean(times, rows[:, column]) for column in range(rows.shape[1])]
    ordinal = None
    if growth is not None and all(value is not None for value in growth):
        lowest = int(np.argmin(growth))
        ordinal = dict(lowest_growth_group=lowest, highest_cost_ratio_group=int(np.argmax(ratios)),
                       satisfied=lowest == int(np.argmax(ratios)))
    return dict(recovery=recovery, departed=departed, maximum_departure_total_variation=largest,
                time_weighted_log_share_growth=growth, ordinal=ordinal)


def development_diagnostics(units, ratios, config):
    """In-sample endpoints on development or mixture vessels (training data)."""
    domain = config['postpulse_domain_days']
    result = {}
    for unit in units:
        keep = (unit['times'] >= domain[0]) & (unit['times'] <= domain[1])
        result[unit['unit_id']] = observed_endpoints(unit['times'][keep], unit['shares'][keep],
                                                     unit['baseline_shares'], ratios, config)
    satisfied = [row['ordinal']['satisfied'] for row in result.values() if row['ordinal']]
    return dict(vessels=result, ordinal_satisfied=int(sum(satisfied)), ordinal_scored=len(satisfied))


def freeze_computation(tables, config, amendment):
    """Every frozen quantity; deterministic, and blind to held-out outcomes."""
    grid = [float(day) for day in config['prediction_grid_days']]
    domain = [float(day) for day in config['postpulse_domain_days']]
    main_records = parse_rows(tables['main'], config)
    held_decoded = [row['source_row'] for row in main_records
                    if row['category'] == config['evaluation_category'] and row['time'] > 0 and row['decoded']]
    if held_decoded:
        raise RuntimeError(f'Held-out post-pulse rows were decoded before the freeze: {held_decoded}')
    main = vessel_map(main_records)
    mixtures = vessel_map(parse_rows(tables['no_herbivore'], config))
    partition = whole_unit_partition([dict(unit_id=key, category=vessel['category']) for key, vessel in main.items()],
                                     config['development_category'], config['evaluation_category'])
    if partition['unsupported']:
        raise ValueError(f'Unsupported vessel categories: {partition["unsupported"]}')
    mixture_ids = sorted(key for key, vessel in mixtures.items() if vessel['treatment'] == MIXTURE_TREATMENT)
    calibration = calibrate(tables['stoichiometry'], config)
    checks = data_checks(main_records, parse_rows(tables['no_herbivore'], config))
    results = []
    for variant in variants(calibration, config):
        rho = np.array(variant['densities'])
        development_units = [unit_trajectory(main[key], rho) for key in partition['development']]
        mixture_units = [unit_trajectory(mixtures[key], rho) for key in mixture_ids]
        missing = [key for key, unit in zip(partition['development'] + mixture_ids,
                                            development_units + mixture_units) if unit is None]
        if missing:
            raise ValueError(f'Training vessels without an eligible baseline: {missing}')
        development = mean_response(development_units, grid, config['minimum_development_contributors'])
        mixture = mean_response(mixture_units, grid, config['minimum_preliminary_mixture_contributors'])
        strengths = fit_strengths(development_units, grid, variant['cost_ratios'],
                                  config['minimum_development_contributors'])
        forecasts = {key: heldout_forecast(main[key], rho, variant['cost_ratios'], development, mixture,
                                           strengths, grid, domain)
                     for key in partition['validation']}
        budget = None
        if variant['resource'] == 'nitrogen':
            budget = {unit['unit_id']: unit['baseline_total'] for unit in development_units + mixture_units}
            budget.update({key: forecast['baseline_total'] for key, forecast in forecasts.items()})
        results.append(dict(variant,
                            development=dict(grid=grid, **development),
                            no_herbivore=dict(grid=grid, **mixture),
                            two_budget=strengths, forecasts=forecasts,
                            development_diagnostics=development_diagnostics(development_units, variant['cost_ratios'], config),
                            no_herbivore_diagnostics=development_diagnostics(mixture_units, variant['cost_ratios'], config),
                            baseline_algal_nitrogen_umol_per_l=budget))
    return jsonable(dict(calibration=calibration, partition=partition, no_herbivore_mixture_vessels=mixture_ids,
                         data_checks=checks, variants=results,
                         budget_reference=dict(inflow_nitrogen_umol_per_l=INFLOW_NITROGEN_UMOL_PER_L,
                                               scope='Reconstructed algal nitrogen at baseline uses pre-pulse preliminary densities; descriptive transfer check, not used by any forecast or score')))


def data_checks(main_records, mixture_records):
    discrepancy, missing = [], []
    for record in main_records + mixture_records:
        if not record['decoded']:
            continue
        if not complete_volumes(record):
            missing.append(dict(chemostat=record['chemostat'], time=record['time'], source_row=record['source_row']))
        elif record['total_algae'] and record['total_algae'] > 0:
            discrepancy.append(abs(sum(record['biovolumes'])-record['total_algae'])/record['total_algae'])
    return dict(decoded_rows_with_incomplete_groups=missing,
                maximum_relative_total_discrepancy=max(discrepancy) if discrepancy else None,
                rows_with_total_discrepancy_above_one_ppm=int(sum(value > 1e-6 for value in discrepancy)),
                scope='Group sums against the supplied total-algae column on decoded rows only')


def score_vessel(forecast, vessel, rho, config, ratios, bounds=None, pooled_index=None):
    times = np.array(forecast['times'], dtype=float)
    rows = {row['time']: row for row in vessel['rows']}
    if not all(rows[time]['decoded'] for time in forecast['times']):
        raise RuntimeError(f'Scoring requires decoded held-out rows: {vessel["id"]}')
    volumes = np.array([[np.nan if value is None else value for value in rows[time]['biovolumes']]
                        for time in forecast['times']])
    observed = resource_profile(volumes, rho)
    models = composition_models(config)
    predicted = {model: as_array(forecast['models'][model]['shares']) for model in models}
    support = common_comparison_support(observed['shares'], predicted)['valid']

    def masked(values):
        result = np.array(values, dtype=float)
        result[~support] = np.nan
        return result

    shares = masked(observed['shares'])
    baseline = np.array(forecast['baseline_shares'])
    composition, recovery, envelope = {}, {}, {}
    for model in models:
        score = trajectory_scores(times, masked(predicted[model]), shares)
        composition[model] = dict(time_weighted_total_variation=score['time_weighted_total_variation'],
                                  total_variation=score['total_variation'],
                                  maximum_absolute_share_error=score['maximum_absolute_share_error'],
                                  covered_duration=score['covered_duration'])
        recovery[model] = recovery_endpoint(times, masked(predicted[model]), baseline,
                                            config['recovery_margin_tv'], 0.,
                                            config['recovery_consecutive_observations'])['status']
        if bounds is not None:
            lower, upper = np.full(len(times), np.nan), np.full(len(times), np.nan)
            for row in np.flatnonzero(support):
                result = profile_tv_envelope(predicted[model][row], volumes[row], rho, pooled_index, bounds)
                lower[row], upper[row] = result['lower'], result['upper']
            envelope[model] = dict(time_weighted_lower=time_weighted_mean(times, lower),
                                   time_weighted_upper=time_weighted_mean(times, upper))
    observed_total = np.where(observed['valid'], observed['total'], np.nan)
    stock = {}
    for model in TOTAL_MODELS:
        total = as_array(forecast['models'][model]['total'])
        error = np.full(len(times), np.nan)
        usable = support & np.isfinite(total) & (total > 0) & np.isfinite(observed_total) & (observed_total > 0)
        error[usable] = np.abs(np.log(total[usable])-np.log(observed_total[usable]))
        stock[model] = dict(time_weighted_absolute_log_error=time_weighted_mean(times, error), absolute_log_error=error)
    projection = {model: (float(np.nanmax(as_array(forecast['models'][model]['projection_distance'])[support]))
                          if np.any(support) else None) for model in ANCHORED_MODELS}
    endpoints = observed_endpoints(times, shares, baseline, ratios, config)
    capacity = forecast['two_budget_capacity_total_variation']
    largest = endpoints['maximum_departure_total_variation']
    endpoints['exceeds_two_budget_capacity'] = None if largest is None else largest > capacity
    return dict(times=times, support=support, observed_shares=shares, observed_total=observed_total,
                composition=composition, total_stock=stock, maximum_projection_distance=projection,
                forecast_recovery=recovery, observed=endpoints, pooled_envelope=envelope or None)


def mean_or_none(values):
    values = [value for value in values if value is not None]
    return float(np.mean(values)) if values else None


def summarize(scores, models):
    summary = {}
    for model in models:
        values = [row['composition'][model]['time_weighted_total_variation'] for row in scores.values()
                  if row['composition'][model]['time_weighted_total_variation'] is not None]
        summary[model] = dict(equal_vessel_mean_time_weighted_total_variation=(float(np.mean(values)) if values else None),
                              scored_vessels=len(values))
    return summary


def verdicts(variant_results, config, amendment):
    rule = amendment['decision_rule']
    margin = rule['practical_margin_tv']
    nitrogen = {row['scenario']: row for row in variant_results if row['resource'] == 'nitrogen'}
    nominal = nitrogen[config['nominal_pooled_scenario']]['vessels']
    result = []
    for first, second in rule['pairs']:
        differences = [nitrogen[scenario]['primary'][second]['equal_vessel_mean_time_weighted_total_variation']
                       - nitrogen[scenario]['primary'][first]['equal_vessel_mean_time_weighted_total_variation']
                       for scenario in config['pooled_density_scenarios']]
        paired = [(row['composition'][first]['time_weighted_total_variation'],
                   row['composition'][second]['time_weighted_total_variation']) for row in nominal.values()]
        paired = [pair for pair in paired if None not in pair]
        first_wins = sum(a < b for a, b in paired)
        second_wins = sum(b < a for a, b in paired)
        if all(value >= margin for value in differences) and first_wins > len(paired)/2:
            verdict = f'{first} outperforms {second}'
        elif all(value <= -margin for value in differences) and second_wins > len(paired)/2:
            verdict = f'{second} outperforms {first}'
        else:
            verdict = 'not distinguished'
        result.append(dict(pair=[first, second], improvement_by_scenario=dict(zip(config['pooled_density_scenarios'], differences)),
                           nominal_vessels_first_lower=first_wins, nominal_vessels_second_lower=second_wins,
                           nominal_scored_vessels=len(paired), verdict=verdict))
    return result


def evaluation_computation(frozen, tables, config, amendment):
    """Score frozen forecasts against decoded held-out observations."""
    main = vessel_map(parse_rows(tables['main'], config))
    models = composition_models(config)
    pooled_index = config['groups'].index('Monoraphidium_Chlorella')
    results = []
    for variant in frozen['variants']:
        rho = np.array(variant['densities'])
        bounds = variant['pooled_bounds'] if variant['scenario'] == config['nominal_pooled_scenario'] else None
        vessels = {key: score_vessel(forecast, main[key], rho, config, variant['cost_ratios'],
                                     bounds, pooled_index)
                   for key, forecast in variant['forecasts'].items()}
        primary = summarize(vessels, models)
        stock = {model: mean_or_none(row['total_stock'][model]['time_weighted_absolute_log_error']
                                     for row in vessels.values()) for model in TOTAL_MODELS}
        envelope = None
        if bounds is not None:
            envelope = {model: dict(equal_vessel_mean_lower=mean_or_none(row['pooled_envelope'][model]['time_weighted_lower'] for row in vessels.values()),
                                    equal_vessel_mean_upper=mean_or_none(row['pooled_envelope'][model]['time_weighted_upper'] for row in vessels.values()))
                        for model in models}
        exceeded = [row['observed']['exceeds_two_budget_capacity'] for row in vessels.values()
                    if row['observed']['exceeds_two_budget_capacity'] is not None]
        ordinal = [row['observed']['ordinal'] for row in vessels.values() if row['observed']['ordinal']]
        departed = [row['observed']['ordinal'] for row in vessels.values()
                    if row['observed']['ordinal'] and row['observed']['departed']]
        agreement = {model: sum(row['forecast_recovery'][model] == row['observed']['recovery']['status']
                                for row in vessels.values()) for model in models}
        results.append(dict(variant_id=variant['variant_id'], resource=variant['resource'], scenario=variant['scenario'],
                            primary=primary,
                            ranking=sorted(models, key=lambda model: primary[model]['equal_vessel_mean_time_weighted_total_variation']),
                            total_stock_equal_vessel_mean_absolute_log_error=stock,
                            ranking_total_stock=sorted(TOTAL_MODELS, key=lambda model: stock[model]),
                            pooled_envelope=envelope,
                            recovery=dict(observed={key: row['observed']['recovery']['status'] for key, row in vessels.items()},
                                          forecast_status_agreement=agreement),
                            two_budget_capacity=dict(scored=len(exceeded), observed_departure_exceeds=int(sum(exceeded))),
                            ordinal=dict(scored=len(ordinal), satisfied=int(sum(item['satisfied'] for item in ordinal)),
                                         departed_scored=len(departed),
                                         departed_satisfied=int(sum(item['satisfied'] for item in departed)),
                                         chance_reference=1./len(config['groups'])),
                            maximum_projection_distance={model: max((row['maximum_projection_distance'][model] for row in vessels.values()
                                                                      if row['maximum_projection_distance'][model] is not None), default=None)
                                                         for model in ANCHORED_MODELS},
                            vessels=vessels))
    return jsonable(dict(variants=results, verdicts=verdicts(results, config, amendment)))


def workbook_paths():
    return {key: WORKBOOKS/name for key, name in INPUTS.items()}


def verify_inputs():
    members = {item['path']: item['sha256'] for item in read_json(SOURCES/'archive-members.json')}
    references = {}
    for key, path in workbook_paths().items():
        reference = ref(path)
        if members.get(reference['path']) != reference['sha256']:
            raise RuntimeError(f'Workbook differs from the retained archive inventory: {reference["path"]}')
        references[key] = reference
    for key, name in (('stoichiometry', 'preliminary-stoichiometry.json'), ('no_herbivore', 'preliminary-no-herbivore.json')):
        if read_json(SOURCES/name)['input_reference'] != references[key]:
            raise RuntimeError(f'Audit copy {name} references different workbook bytes')
    for name in ('acquisition.json', 'archive-members.json', 'workbook-schema.json',
                 'preliminary-stoichiometry.json', 'preliminary-no-herbivore.json', 'source-audit.json'):
        references[name] = ref(SOURCES/name)
    return references


def crosscheck(tables, config):
    """Compare this reader with the retained openpyxl audit copies (non-held-out cells)."""
    checked = 0
    copy = {row['source_row']: row for row in read_json(SOURCES/'preliminary-stoichiometry.json')['rows']}
    for row in tables['stoichiometry']['rows']:
        expected = copy[row['source_row']]
        for header in STOICHIOMETRY_META[1:]:
            if row['metadata'][header] != expected[header]:
                raise RuntimeError(f'Reader mismatch: stoichiometry row {row["source_row"]} {header}')
        if serial_date(row['metadata']['Date'], tables['stoichiometry']['date1904']) != expected['Date']:
            raise RuntimeError(f'Reader date mismatch: stoichiometry row {row["source_row"]}')
        for header in STOICHIOMETRY_VALUES:
            if row['values'][header] != expected[header]:
                raise RuntimeError(f'Reader mismatch: stoichiometry row {row["source_row"]} {header}')
            checked += 1
    if len(copy) != len(tables['stoichiometry']['rows']):
        raise RuntimeError('Stoichiometry row count differs from the audit copy')
    copy = {row['source_row']: row for row in read_json(SOURCES/'preliminary-no-herbivore.json')['rows']}
    for row in tables['no_herbivore']['rows']:
        expected = copy[row['source_row']]
        for header in META:
            value = row['metadata'][header]
            if header == 'Date':
                value = serial_date(value, tables['no_herbivore']['date1904'])
            if value != expected[header]:
                raise RuntimeError(f'Reader mismatch: no-herbivore row {row["source_row"]} {header}')
        for header in [TOTAL_HEADER] + group_headers(config):
            if row['values'][header] != expected[header]:
                raise RuntimeError(f'Reader mismatch: no-herbivore row {row["source_row"]} {header}')
            checked += 1
    if len(copy) != len(tables['no_herbivore']['rows']):
        raise RuntimeError('No-herbivore row count differs from the audit copy')
    schema = read_json(SOURCES/'workbook-schema.json')['workbooks'][0]['sheets'][0]['metadata_rows']
    expected_rows = {row['source_row']: row for row in schema}
    for row in tables['main']['rows']:
        expected = expected_rows[row['source_row']]
        for header in META:
            value = row['metadata'][header]
            if header == 'Date':
                value = serial_date(value, tables['main']['date1904'])
            if value != expected[header]:
                raise RuntimeError(f'Reader mismatch: main row {row["source_row"]} {header}')
            checked += 1
    if len(expected_rows) != len(tables['main']['rows']):
        raise RuntimeError('Main row count differs from the metadata audit')
    return dict(cells_matched=checked, compared_with='openpyxl audit copies of calibration tables and main metadata')


def decoded_records(tables, config):
    return dict(stoichiometry=[dict(source_row=row['source_row'], metadata=row['metadata'], values=row['values'])
                               for row in tables['stoichiometry']['rows']],
                no_herbivore=parse_rows(tables['no_herbivore'], config),
                main=[row for row in parse_rows(tables['main'], config) if row['decoded']])


def freeze():
    target = DATA/'frozen-forecasts.json'
    if target.exists():
        verify_freeze()
        return
    if (RESULTS/'study.json').exists() or (DATA/'evaluation-records.json').exists():
        raise RuntimeError('Cannot freeze after evaluation outputs exist')
    config, amendment = read_json(CONFIG), read_json(AMENDMENTS[0])
    if any(read_json(path)['run_id'] != config['run_id'] for path in AMENDMENTS):
        raise RuntimeError('Amendment belongs to a different run')
    inputs = verify_inputs()
    tables = read_tables(workbook_paths(), config, 'freeze')
    reader = crosscheck(tables, config)
    computation = freeze_computation(tables, config, amendment)
    sources, snapshots = [], []
    for source in SOURCE_PATHS:
        snapshot = DATA/'original-sources'/source
        snapshot.parent.mkdir(parents=True, exist_ok=True)
        if not snapshot.exists():
            snapshot.write_bytes((ROOT/source).read_bytes())
        sources.append(ref(ROOT/source))
        snapshots.append(ref(snapshot))
        if sources[-1]['sha256'] != snapshots[-1]['sha256']:
            raise RuntimeError(f'Source snapshot differs: {source}')
    save_json(DATA/'development-records.json', jsonable(decoded_records(tables, config)))
    save_json(target, dict(
        run_id=config['run_id'], status='frozen_before_heldout_decoding', frozen_utc=now(),
        held_out_postpulse_values_decoded=False,
        gated_heldout_rows=tables['main']['gated_rows'],
        config_reference=ref(CONFIG), amendment_references=[ref(path) for path in AMENDMENTS],
        input_references=inputs,
        source_references=sources, source_snapshot_references=snapshots,
        development_records_reference=ref(DATA/'development-records.json'),
        reader_crosscheck=reader,
        git_head_at_freeze=subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        environment=dict(python=sys.version, numpy=np.__version__, platform=platform.platform()),
        scope=config['analysis_separation'], computation=computation))


def verify_freeze():
    frozen = read_json(DATA/'frozen-forecasts.json')
    references = [frozen['config_reference'], *frozen['amendment_references'],
                  *frozen['input_references'].values(), *frozen['source_references'],
                  *frozen['source_snapshot_references'], frozen['development_records_reference']]
    for reference in references:
        verify_reference(reference)
    return frozen


def replay_freeze(frozen):
    config, amendment = read_json(CONFIG), read_json(AMENDMENTS[0])
    tables = read_tables(workbook_paths(), config, 'freeze')
    if freeze_computation(tables, config, amendment) != frozen['computation']:
        raise RuntimeError('Freeze replay differs from the retained forecasts')
    if crosscheck(tables, config) != frozen['reader_crosscheck']:
        raise RuntimeError('Reader cross-check replay differs')
    return config, amendment


def heldout_crosscheck(tables, config):
    """After the freeze: compare held-out decoded cells with openpyxl."""
    import openpyxl
    workbook = openpyxl.load_workbook(workbook_paths()['main'], read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    columns = [header.index(name) for name in [TOTAL_HEADER] + group_headers(config)]
    expected = {number: row for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)}
    workbook.close()
    checked = 0
    for row in tables['main']['rows']:
        if row['metadata']['Category'] != config['evaluation_category'] or row['metadata']['Time standardised to pulse'] <= 0:
            continue
        for name, column in zip([TOTAL_HEADER] + group_headers(config), columns):
            if row['values'][name] != expected[row['source_row']][column]:
                raise RuntimeError(f'Held-out reader mismatch at row {row["source_row"]}: {name}')
            checked += 1
    return dict(cells_matched=checked, compared_with='openpyxl cached values of held-out post-pulse cells')


def evaluate():
    if (RESULTS/'study.json').exists():
        audit()
        return
    frozen = verify_freeze()
    config, amendment = replay_freeze(frozen)
    started = now()
    if started <= frozen['frozen_utc']:
        raise RuntimeError('Evaluation would precede the freeze')
    tables = read_tables(workbook_paths(), config, 'evaluation')
    checked = heldout_crosscheck(tables, config)
    records = [row for row in parse_rows(tables['main'], config)
               if row['category'] == config['evaluation_category'] and row['time'] > 0]
    save_json(DATA/'evaluation-records.json', jsonable(records))
    evaluation = evaluation_computation(frozen['computation'], tables, config, amendment)
    report = dict(run_id=config['run_id'], status='complete', kind='retrospective_chemostat_resource_response',
                  evaluation_started_utc=started, frozen_forecasts=ref(DATA/'frozen-forecasts.json'),
                  config_reference=frozen['config_reference'], amendment_references=frozen['amendment_references'],
                  evaluation_records_reference=ref(DATA/'evaluation-records.json'),
                  heldout_reader_crosscheck=checked, git_head_at_evaluation=subprocess.check_output(
                      ['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
                  evaluation=evaluation,
                  interpretation=('Separately calibrated preliminary-culture resource densities convert observed group biovolumes into resource-stock proxies; '
                                  'forecasts transfer development responses or a development-fitted two-budget cost-ratio rule to whole held-out vessels. '
                                  'Three partly pooled groups do not test logarithmic neutrality or a power-law exponent.'),
                  statistical_scope=config['uncertainty_scope'])
    save_json(RESULTS/'study.json', report)
    save_json(RESULTS/'output-manifest.json', dict(run_id=config['run_id'], artifacts=[ref(RESULTS/'study.json')],
                                                   evaluation_records=ref(DATA/'evaluation-records.json')))


def audit():
    frozen = verify_freeze()
    config, amendment = replay_freeze(frozen)
    report = read_json(RESULTS/'study.json')
    for key in ('frozen_forecasts', 'evaluation_records_reference'):
        verify_reference(report[key])
    manifest = read_json(RESULTS/'output-manifest.json')
    for reference in manifest['artifacts'] + [manifest['evaluation_records']]:
        verify_reference(reference)
    tables = read_tables(workbook_paths(), config, 'evaluation')
    if heldout_crosscheck(tables, config) != report['heldout_reader_crosscheck']:
        raise RuntimeError('Held-out reader cross-check replay differs')
    if evaluation_computation(frozen['computation'], tables, config, amendment) != report['evaluation']:
        raise RuntimeError('Evaluation replay differs from the retained report')
    return dict(status='verified_offline', run_id=config['run_id'])


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--stage', choices=['freeze', 'evaluate', 'audit'], required=True)
    stage = parser.parse_args().stage
    {'freeze': freeze, 'evaluate': evaluate, 'audit': audit}[stage]()
    print(json.dumps(dict(stage=stage, status='complete', run_id=read_json(CONFIG)['run_id'])))


if __name__ == '__main__':
    main()
