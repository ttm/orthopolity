"""Acquire immutable public chemostat sources, or verify them offline.

Run --stage acquire explicitly for first retrieval. Existing files must match
their retained receipt hashes. Archive extraction preserves original bytes and
rejects traversal, links, collisions and unreasonable expanded sizes.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import ssl
import stat
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path, PurePosixPath

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data/chemostat-sources/2026-10-02"
DRYAD_API = "https://datadryad.org/api/v2/datasets/doi%3A10.5061%2Fdryad.51c59zwj5"
FIGSHARE_API = "https://api.figshare.com/v2/collections/8172360"


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_bytes(value: object) -> bytes:
    return (json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n").encode()


def write_exact(path: Path, body: bytes) -> None:
    if path.exists():
        if path.read_bytes() != body:
            raise RuntimeError(f"Refusing to replace changed snapshot: {path.relative_to(ROOT)}")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(body)


def tls_context() -> tuple[ssl.SSLContext, dict]:
    cert = Path("/etc/ssl/cert.pem")
    context = ssl.create_default_context(cafile=str(cert) if cert.exists() else None)
    return context, {
        "certificate_verification": "CERT_REQUIRED",
        "hostname_verification": True,
        "openssl_version": ssl.OPENSSL_VERSION,
        "cafile": str(cert) if cert.exists() else "Python default trust store",
        "cafile_sha256": sha(cert.read_bytes()) if cert.exists() else None,
    }


def acquire() -> dict:
    BASE.mkdir(parents=True, exist_ok=True)
    receipt_path = BASE / "acquisition.json"
    receipt = json.loads(receipt_path.read_text()) if receipt_path.exists() else {
        "schema_version": 1,
        "source_doi": "10.5061/dryad.51c59zwj5",
        "supplement_collection_doi": "10.6084/m9.figshare.c.8172360",
        "started_utc": now(),
        "scope": "Published files retained unchanged; main outcomes not inspected during source audit.",
        "files": [],
    }
    if receipt.get("complete"):
        verify(receipt)
        return receipt
    context, tls = tls_context()
    receipt["tls"] = tls
    inventory = {item["path"]: item for item in receipt["files"]}

    def download(relative: str, url: str, provider_md5: str | None = None,
                 provider_sha256: str | None = None) -> bytes:
        path = BASE / relative
        key = str(path.relative_to(ROOT))
        if key in inventory:
            body = path.read_bytes()
            if sha(body) != inventory[key]["sha256"]:
                raise RuntimeError(f"Changed retained source: {key}")
            return body
        if path.exists():
            raise RuntimeError(f"Unreceipted existing source: {key}")
        started = now()
        request = urllib.request.Request(url, headers={"User-Agent": "OrthopolityResearch/1.0 (public research source audit)"})
        try:
            with urllib.request.urlopen(request, context=context, timeout=45) as response:
                chunks = []
                total = 0
                while chunk := response.read(1024 * 1024):
                    chunks.append(chunk)
                    total += len(chunk)
                    if total > 150_000_000:
                        raise RuntimeError("Source exceeds 150 MB limit")
                    if total % (16 * 1024 * 1024) == 0:
                        print(f"Receiving {relative}: {total} bytes", flush=True)
                body = b"".join(chunks)
                metadata = {
                    "path": key, "requested_url": url, "final_url": response.url,
                    "retrieval_started_utc": started, "retrieval_completed_utc": now(),
                    "http_status": response.status, "headers": dict(response.headers),
                    "bytes": len(body), "sha256": sha(body), "provider_md5": provider_md5,
                    "provider_sha256": provider_sha256,
                }
        except Exception as error:
            attempt = {"requested_url": url, "started_utc": started,
                       "completed_utc": now(), "error": f"{type(error).__name__}: {error}"}
            with (BASE / "retrieval-attempts.jsonl").open("a") as stream:
                stream.write(json.dumps(attempt, sort_keys=True) + "\n")
            raise
        if provider_md5 and hashlib.md5(body).hexdigest() != provider_md5:
            raise RuntimeError(f"Provider MD5 mismatch: {key}")
        if provider_sha256 and sha(body) != provider_sha256:
            raise RuntimeError(f"Provider SHA-256 mismatch: {key}")
        write_exact(path, body)
        receipt["files"].append(metadata)
        inventory[key] = metadata
        receipt_path.write_bytes(json_bytes(receipt))
        print(f"Retained {relative}: {len(body)} bytes", flush=True)
        return body

    dataset = json.loads(download("dryad/dataset-metadata.json", DRYAD_API))
    links = dataset.get("_links", {})
    version_link = links.get("stash:version", links.get("stash:latestVersion"))
    if version_link:
        version_url = version_link["href"]
        if version_url.startswith("/"):
            version_url = "https://datadryad.org" + version_url
        version = json.loads(download("dryad/version-metadata.json", version_url))
    else:
        version = dataset
    files_link = version.get("_links", {}).get("stash:files") or links.get("stash:files")
    if not files_link:
        raise RuntimeError(f"No Dryad files link; retained metadata keys: {list(version)}")
    files_url = files_link["href"]
    if files_url.startswith("/"):
        files_url = "https://datadryad.org" + files_url
    files = json.loads(download("dryad/files-metadata.json", files_url))
    entries = files.get("_embedded", {}).get("stash:files", files.get("files", []))
    if not entries:
        raise RuntimeError(f"No Dryad file entries: {list(files)}")
    # Public supplementary deposits are enumerated independently, so their
    # methods remain available if the main archive download is unavailable.
    collection = json.loads(download("figshare/collection-metadata.json", FIGSHARE_API))
    articles = json.loads(download("figshare/articles-metadata.json", FIGSHARE_API + "/articles"))
    for article in articles:
        article_url = article.get("url") or f"https://api.figshare.com/v2/articles/{article['id']}"
        detail = json.loads(download(f"figshare/article-{article['id']}-metadata.json", article_url))
        for item in detail.get("files", []):
            name = PurePosixPath(item["name"]).name
            if name.lower().endswith(".pdf"):
                download(f"figshare/{item['id']}-{name}", item["download_url"], item.get("supplied_md5") or None)
    for record_id in (14773027, 14773023):
        download(f"zenodo/record-{record_id}-metadata.json", f"https://zenodo.org/api/records/{record_id}")
    unavailable = []
    for entry in entries:
        name = entry.get("path", entry.get("name"))
        if name not in ("data.zip", "README.md"):
            continue
        # The API download route needs authentication; the dataset's published
        # web links provide the same public files without credentials.
        self_link = entry.get("_links", {}).get("self", {}).get("href", "")
        file_id = self_link.rsplit("/", 1)[-1]
        file_url = f"https://datadryad.org/downloads/file_stream/{file_id}" if file_id.isdecimal() else entry.get("download_url")
        if not file_url:
            raise RuntimeError(f"No download URL for {name}: {entry}")
        if file_url.startswith("/"):
            file_url = "https://datadryad.org" + file_url
        try:
            download(f"dryad/{name}", file_url,
                     entry.get("digest") if entry.get("digestType") == "md5" else None,
                     entry.get("digest") if entry.get("digestType") == "sha-256" else None)
        except urllib.error.HTTPError as error:
            if error.code not in (401, 403):
                raise
            unavailable.append({"name": name, "requested_url": file_url,
                                "http_status": error.code, "provider_digest": entry["digest"],
                                "provider_digest_type": entry["digestType"]})
    if not (BASE / "dryad/data.zip").exists():
        mirror = json.loads((BASE / "zenodo/record-14773023-metadata.json").read_text())
        for entry in mirror["files"]:
            if entry["key"] == "scripts_data_plankton_responses_Npulse.zip":
                checksum = entry["checksum"]
                if not checksum.startswith("md5:"):
                    raise RuntimeError("Unrecognized Zenodo checksum")
                download("zenodo/scripts_data_plankton_responses_Npulse.zip",
                         entry["links"]["self"], checksum[4:])
    if not (BASE / "dryad/data.zip").exists() and not (BASE / "zenodo/scripts_data_plankton_responses_Npulse.zip").exists():
        raise RuntimeError("Neither main Dryad archive nor author Zenodo bundle available")

    if not list((BASE / "figshare").glob("*.pdf")):
        raise RuntimeError("No supplementary PDF found")
    receipt["source_publication"] = dataset.get("relatedWorks", [])
    receipt["unavailable_dryad_files"] = unavailable
    receipt["source_selection"] = "dryad_original" if (BASE / "dryad/data.zip").exists() else "author_public_zenodo_bundle"
    receipt["complete"] = True
    receipt["completed_utc"] = now()
    receipt_path.write_bytes(json_bytes(receipt))
    verify(receipt)
    return receipt


def extract_archive(path: Path) -> list[dict]:
    expanded = path.parent / "extracted"
    inventory = []
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if sum(item.file_size for item in entries) > 200_000_000:
            raise RuntimeError("Archive expansion exceeds 200 MB")
        names = set()
        for item in entries:
            name = PurePosixPath(item.filename)
            mode = item.external_attr >> 16
            if name.is_absolute() or ".." in name.parts or "\\" in item.filename or stat.S_ISLNK(mode):
                raise RuntimeError(f"Unsafe archive member: {item.filename}")
            if item.filename in names:
                raise RuntimeError(f"Duplicate archive member: {item.filename}")
            names.add(item.filename)
            if item.is_dir():
                continue
            destination = expanded.joinpath(*name.parts)
            body = archive.read(item)
            write_exact(destination, body)
            inventory.append({"archive_member": item.filename, "path": str(destination.relative_to(ROOT)),
                              "bytes": len(body), "sha256": sha(body)})
    write_exact(BASE / "archive-members.json", json_bytes(inventory))
    return inventory


def verify(receipt: dict | None = None) -> dict:
    if receipt is None:
        receipt = json.loads((BASE / "acquisition.json").read_text())
    for item in receipt["files"]:
        body = (ROOT / item["path"]).read_bytes()
        if len(body) != item["bytes"] or sha(body) != item["sha256"]:
            raise RuntimeError(f"Retained source mismatch: {item['path']}")
        if item.get("provider_md5") and hashlib.md5(body).hexdigest() != item["provider_md5"]:
            raise RuntimeError(f"Retained provider hash mismatch: {item['path']}")
        if item.get("provider_sha256") and sha(body) != item["provider_sha256"]:
            raise RuntimeError(f"Retained provider SHA-256 mismatch: {item['path']}")
    archive = BASE / "dryad/data.zip"
    if not archive.exists():
        archive = BASE / "zenodo/scripts_data_plankton_responses_Npulse.zip"
    members = extract_archive(archive)
    return {"retained_source_files": len(receipt["files"]), "archive_members": len(members),
            "complete": receipt.get("complete", False), "source_selection": receipt.get("source_selection"),
            "unavailable_dryad_files": [item["name"] for item in receipt.get("unavailable_dryad_files", [])]}


def serializable(value: object) -> object:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Nonfinite workbook value")
    return value


def audit() -> dict:
    """Inspect calibration values and main metadata; do not inspect outcomes.

    Main iter_rows is restricted to the five identifier/time columns. Outcome
    missingness uses XML element presence, never parses numeric outcome values.
    Reading original archive bytes is not a claim of global outcome blinding.
    """
    import openpyxl
    verify()
    candidates = list(BASE.glob("*/extracted/**/Chemostat_experimental_timeseries.xlsx"))
    if len(candidates) != 1:
        raise RuntimeError(f"Expected one main workbook, found {len(candidates)}")
    main = candidates[0]
    folder = main.parent
    calibration_names = (
        "algae_stoichiometry_preliminary_experiments.xlsx",
        "algae_no-rotifer_chemostat_preliminary_experiments.xlsx",
    )
    schema = {
        "inspection_scope": "Main headers, identifiers, dates, times and cell-presence missingness; preliminary values only.",
        "main_outcome_values_inspected": False,
        "workbooks": [],
    }
    references = []
    for path in (main, *(folder / name for name in calibration_names)):
        reference = {"path": str(path.relative_to(ROOT)), "sha256": sha(path.read_bytes())}
        references.append(reference)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        item = {"input_reference": reference, "sheets": []}
        for sheet in workbook.worksheets:
            header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            sheet_info = {"sheet": sheet.title, "row_count_including_header": sheet.max_row,
                          "column_count": sheet.max_column, "headers": header}
            if path == main:
                metadata_rows = []
                vessel_rows = defaultdict(list)
                for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=5, values_only=True), 2):
                    if all(value is None for value in row):
                        continue
                    record = dict(zip(header[:5], (serializable(value) for value in row)))
                    record["source_row"] = row_number
                    metadata_rows.append(record)
                    vessel_rows[str(row[2])].append(record)
                summaries = []
                for vessel, rows in sorted(vessel_rows.items()):
                    times = [row[header[4]] for row in rows]
                    if any(not isinstance(t, (int, float)) for t in times):
                        raise RuntimeError(f"Non-numeric source time in vessel {vessel}")
                    if len(times) != len(set(times)):
                        raise RuntimeError(f"Duplicate source times in vessel {vessel}")
                    summaries.append({
                        "chemostat": vessel, "category": sorted({row[header[0]] for row in rows}),
                        "treatment": sorted({row[header[1]] for row in rows}),
                        "n_rows": len(rows), "pre_pulse_times": sorted(t for t in times if t < 0),
                        "pulse_time_available": 0 in times,
                        "post_pulse_times": sorted(t for t in times if t > 0),
                        "date_range": [min(row[header[3]] for row in rows), max(row[header[3]] for row in rows)],
                    })
                sheet_info["metadata_rows"] = metadata_rows
                sheet_info["vessels"] = summaries
                sheet_info["row_counts_by_treatment"] = dict(sorted(Counter(row[header[1]] for row in metadata_rows).items()))
                # Presence inspection reads no numeric outcome cell content.
                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                missing = defaultdict(Counter)
                phase_by_row = {r["source_row"]: "pre_or_at_pulse" if r[header[4]] <= 0 else "post_pulse" for r in metadata_rows}
                with zipfile.ZipFile(path) as archive:
                    for event, row_element in ET.iterparse(archive.open("xl/worksheets/sheet1.xml"), events=("end",)):
                        if row_element.tag != ns + "row":
                            continue
                        row_number = int(row_element.attrib["r"])
                        phase = phase_by_row.get(row_number)
                        if phase:
                            cells = {cell.attrib["r"].rstrip("0123456789"): cell for cell in row_element.findall(ns + "c")}
                            for column_number in range(6, len(header) + 1):
                                letter = openpyxl.utils.get_column_letter(column_number)
                                cell = cells.get(letter)
                                present = cell is not None and (cell.find(ns + "v") is not None or cell.find(ns + "is") is not None)
                                missing[phase][header[column_number - 1]] += not present
                        row_element.clear()
                sheet_info["absent_value_elements_by_phase"] = {phase: dict(counts) for phase, counts in sorted(missing.items())}
                sheet_info["missingness_scope"] = "Counts omitted/empty value elements only; textual NA and numeric validity are not evaluated for main outcomes."
            else:
                raw_rows = []
                for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    if all(value is None for value in row):
                        continue
                    record = {str(key): serializable(value) for key, value in zip(header, row) if key is not None}
                    record["source_row"] = row_number
                    raw_rows.append(record)
                prepared = {
                    "input_reference": reference, "source_sheet": sheet.title,
                    "values_are": "Original cached workbook values, unchanged units; blank values remain null.",
                    "columns": [column for column in header if column is not None],
                    "rows": raw_rows,
                }
                output_name = "preliminary-stoichiometry.json" if "stoichiometry" in path.name else "preliminary-no-herbivore.json"
                write_exact(BASE / output_name, json_bytes(prepared))
                sheet_info["retained_preliminary_json"] = str((BASE / output_name).relative_to(ROOT))
                sheet_info["n_nonempty_records"] = len(raw_rows)
                sheet_info["missing_values_by_column"] = {str(column): sum(record[str(column)] is None for record in raw_rows) for column in header if column is not None}
                if "stoichiometry" in path.name:
                    sheet_info["algae"] = sorted({record["Algae"] for record in raw_rows})
                    sheet_info["timepoints"] = sorted({record["Time standardised to pulse [days]"] for record in raw_rows})
                    sheet_info["replicate_identifier_present"] = any("rep" in str(column).lower() for column in header if column)
            item["sheets"].append(sheet_info)
        workbook.close()
        schema["workbooks"].append(item)
    write_exact(BASE / "workbook-schema.json", json_bytes(schema))
    audit_summary = {
        "schema_version": 1,
        "input_references": references + [
            {"path": str((BASE / "acquisition.json").relative_to(ROOT)), "sha256": sha((BASE / "acquisition.json").read_bytes())},
            {"path": str((BASE / "archive-members.json").relative_to(ROOT)), "sha256": sha((BASE / "archive-members.json").read_bytes())},
        ],
        "source_selection": json.loads((BASE / "acquisition.json").read_text())["source_selection"],
        "main_workbook": references[0],
        "main_vessels": len(schema["workbooks"][0]["sheets"][0]["vessels"]),
        "main_rows": len(schema["workbooks"][0]["sheets"][0]["metadata_rows"]),
        "preliminary_stoichiometry_rows": schema["workbooks"][1]["sheets"][0]["n_nonempty_records"],
        "preliminary_replicate_identifier_present": schema["workbooks"][1]["sheets"][0]["replicate_identifier_present"],
        "main_outcome_values_inspected": False,
        "full_data_fitted_posteriors_used": False,
        "software": {"python": __import__("sys").version, "openpyxl": openpyxl.__version__},
        "auditor_source_reference": {"path": str(Path(__file__).resolve().relative_to(ROOT)), "sha256": sha(Path(__file__).read_bytes())},
    }
    write_exact(BASE / "source-audit.json", json_bytes(audit_summary))
    return {key: audit_summary[key] for key in ("main_vessels", "main_rows", "preliminary_stoichiometry_rows", "preliminary_replicate_identifier_present", "main_outcome_values_inspected")}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", choices=("acquire", "verify", "audit"), default="verify")
    args = parser.parse_args()
    if args.stage == "acquire":
        acquire()
    print(json.dumps(audit() if args.stage == "audit" else verify(), indent=2, sort_keys=True))
