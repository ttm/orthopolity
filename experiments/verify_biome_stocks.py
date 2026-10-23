"""Independently verify the released BioMe diagnostic, without importing its code.

Uses openpyxl against raw OOXML for every retained absorbance cell, explicit
Python loops for stocks, and centered covariance formulas for calibration.
Run only after the diagnostic's committed freeze has been pushed and released.
"""
import argparse
import hashlib
import json
import math
from pathlib import Path
import statistics
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data/physics-candidates/2026-10-23/biome'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
LEVELS = [450 - 50 * i for i in range(10)]
PORES = [0., .03, .1, .2, .4]
DYES = ['phenol_red', 'bromocresol_purple']


def load_checked_blocks(path, starts, count):
    """Check all retained cells with two independent workbook decoders."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if book.sheetnames != ['Plate 1 - Sheet1']:
        raise AssertionError('Unexpected sheet map')
    sheet = book[book.sheetnames[0]]
    with ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    cells = {cell.attrib['r']: cell for cell in xml.findall('s:sheetData/s:row/s:c', NS)}
    blocks, checked, time_checked, maximum_time_difference_seconds = [], 0, 0, 0.
    inventories = []
    for dye, first_row in enumerate(starts):
        first_column = 4 + 30 * dye
        labels = [cell.value for row in sheet.iter_rows(
            min_row=first_row-1, max_row=first_row-1,
            min_col=4, max_col=63) for cell in row]
        expected = [f'{row}{col}' for row in 'ABCDEF' for col in range(1, 11)]
        if labels != expected:
            raise AssertionError('Unexpected plate labels')
        times = []
        for row in sheet.iter_rows(min_row=first_row, max_row=first_row+count-1,
                                   min_col=2, max_col=2):
            cell = row[0]
            raw = float(cells[cell.coordinate].find('s:v', NS).text)
            # openpyxl rounds formatted Excel times to the nearest millisecond.
            decoded = openpyxl.utils.datetime.to_excel(cell.value, book.epoch)
            difference = abs(raw-decoded)*86400
            if difference > .00051:
                raise AssertionError(f'Time decoders disagree: {cell.coordinate}')
            maximum_time_difference_seconds = max(maximum_time_difference_seconds, difference)
            times.append(raw*24)
            time_checked += 1
        if any(not math.isclose(t-times[0], i/4, rel_tol=0, abs_tol=1e-9)
               for i, t in enumerate(times)):
            raise AssertionError('Unexpected time spacing')
        inventories.append({'rows': [first_row, first_row+count-1], 'times': count,
                            'instrument_hours': [times[0], times[-1]],
                            'elapsed_hours': [0., (count-1)/4], 'plate_labels': labels})
        block = []
        for row in sheet.iter_rows(min_row=first_row, max_row=first_row+count-1,
                                   min_col=first_column, max_col=first_column+29):
            readings = []
            for cell in row:
                raw = cells[cell.coordinate]
                if cell.data_type != 'n' or raw.find('s:f', NS) is not None:
                    raise AssertionError(f'Nonoriginal numeric cell: {cell.coordinate}')
                value = float(cell.value)
                if not math.isfinite(value) or value != float(raw.find('s:v', NS).text):
                    raise AssertionError(f'Decoders disagree: {cell.coordinate}')
                readings.append(value)
                checked += 1
            block.append([readings[10*r:10*r+10] for r in range(3)])
        blocks.append(block)
    book.close()
    return blocks, checked, inventories, time_checked, maximum_time_difference_seconds


def linear_calibration(x, y):
    xbar, ybar = statistics.mean(x), statistics.mean(y)
    slope = math.fsum((a-xbar)*(b-ybar) for a, b in zip(x, y)) / math.fsum(
        (a-xbar)**2 for a in x)
    if slope <= 0:
        raise AssertionError('Nonpositive independent calibration slope')
    return ybar - slope*xbar, slope


def calibration(block):
    means = [[statistics.mean(frame[r][c] for frame in block)
              for c in range(10)] for r in range(3)]
    concentrations = LEVELS * 3
    absorbances = [value for row in means for value in row]
    intercept, slope = linear_calibration(concentrations, absorbances)
    held_errors = []
    for held in LEVELS:
        train = [(c, a) for c, a in zip(concentrations, absorbances) if c != held]
        aa, bb = linear_calibration([c for c, _ in train], [a for _, a in train])
        held_errors.extend((a-aa)/bb-held for c, a in zip(concentrations, absorbances)
                           if c == held)
    return {'intercept': intercept, 'slope': slope,
            'rmse_uM': math.sqrt(statistics.mean(
                ((a-intercept)/slope-c)**2 for a, c in zip(absorbances, concentrations))),
            'maximum_held_concentration_error_uM': max(map(abs, held_errors)),
            'maximum_temporal_deviation_uM': max(
                abs(frame[r][c]-means[r][c])/slope
                for frame in block for r in range(3) for c in range(10)),
            'independent_wells': 30, 'repeated_reads_per_well': 25}


def summarize(values):
    valid = [value for value in values if value is not None]
    return {'replicates': values, 'valid_replicates': len(valid),
            'mean': statistics.mean(valid) if valid else None,
            'minimum': min(valid) if valid else None,
            'maximum': max(valid) if valid else None}


def diagnostic(standards, readings):
    cal = calibration(standards)
    concentrations = [[[(value-cal['intercept'])/cal['slope'] for value in row]
                        for row in frame] for frame in readings]
    totals = [[[(row[2*p]+row[2*p+1])/400 for p in range(5)]
               for row in frame] for frame in concentrations]
    if any(value <= 0 for row in totals[0] for value in row):
        raise AssertionError('Nonpositive initial total')
    relative = [[[frame[r][p]/totals[0][r][p] for p in range(5)]
                 for r in range(3)] for frame in totals]
    controls = [[[None for _ in range(5)] for _ in range(3)] for _ in readings]
    invalid = []
    for r in range(3):
        contrasts = [frame[r][0]-frame[r][1] for frame in readings]
        if any(contrast <= 0 for contrast in contrasts):
            invalid.append({'replicate': r+1,
                            'reason': 'nonpositive no-pore control contrast'})
            continue
        for p in range(5):
            values = [(frame[r][2*p]+frame[r][2*p+1]-2*frame[r][1])/contrast
                      for frame, contrast in zip(readings, contrasts)]
            if values[0] <= 0:
                invalid.append({'replicate': r+1, 'pore_um': PORES[p],
                                'reason': 'nonpositive initial control ratio'})
            else:
                for t, value in enumerate(values):
                    controls[t][r][p] = value/values[0]
    traces = {'recovery_nominal': totals, 'recovery_initial': relative,
              'control_ratio_initial': controls}
    groups = []
    for p, pore in enumerate(PORES):
        group = {'pore_um': pore}
        for name, trace in traces.items():
            initial = [trace[0][r][p] for r in range(3)]
            final = [statistics.mean(frame[r][p] for frame in trace[-5:])
                     if trace[-1][r][p] is not None else None for r in range(3)]
            group[name] = {'initial': summarize(initial), 'last_hour': summarize(final)}
        groups.append(group)
    changes = {name: summarize([
        statistics.mean(frame[r][i] for frame in concentrations[-5:])-concentrations[0][r][i]
        for r in range(3)]) for i, name in enumerate(('donor', 'recipient'))}
    return {'calibration': cal, 'groups': groups, 'invalid_controls': invalid,
            'no_pore_changes_uM': changes, 'traces': traces}


def compare(actual, expected, stats, path='diagnostic'):
    if isinstance(actual, dict):
        if not isinstance(expected, dict) or actual.keys() != expected.keys():
            raise AssertionError(f'Different fields at {path}')
        for key, value in actual.items():
            compare(value, expected[key], stats, f'{path}.{key}')
    elif isinstance(actual, list):
        if not isinstance(expected, list) or len(actual) != len(expected):
            raise AssertionError(f'Different list length at {path}')
        for i, (value, wanted) in enumerate(zip(actual, expected)):
            compare(value, wanted, stats, f'{path}[{i}]')
    elif isinstance(actual, (float, int)):
        if not isinstance(expected, (float, int)) or not math.isclose(
                actual, expected, rel_tol=1e-11, abs_tol=1e-10):
            raise AssertionError(f'Independent calculation differs at {path}')
        stats['numerical_comparisons'] += 1
        stats['maximum_absolute_difference'] = max(
            stats['maximum_absolute_difference'], abs(actual-expected))
    elif actual != expected:
        raise AssertionError(f'Values differ at {path}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--diagnostic', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    released = json.loads(args.diagnostic.read_text())
    receipt = json.loads((SOURCE/'acquisition.json').read_text())
    for entry in receipt['files']:
        raw = (SOURCE/entry['file']).read_bytes()
        if len(raw) != entry['bytes'] or hashlib.sha256(raw).hexdigest() != entry['sha256']:
            raise AssertionError(f'Source bytes differ: {entry["file"]}')
    standards, n_standards, standard_inventory, nt_standards, dt_standards = load_checked_blocks(
        SOURCE/'optical-calibration.xlsx', (28, 57), 25)
    responses, n_responses, response_inventory, nt_responses, dt_responses = load_checked_blocks(
        SOURCE/'dye-diffusion.xlsx', (28, 321), 289)
    stats = {'numerical_comparisons': 0, 'maximum_absolute_difference': 0.}
    compare([i/4 for i in range(289)], released['hours'], stats, 'hours')
    compare({'optical-calibration.xlsx': standard_inventory,
             'dye-diffusion.xlsx': response_inventory},
            released['source_inventory']['workbooks'], stats, 'source_inventory.workbooks')
    for dye, name in enumerate(DYES):
        compare(diagnostic(standards[dye], responses[dye]), released['dyes'][name], stats, name)
    result = {'passed': True, 'scope': 'independent verification of the released stock diagnostic',
              'freeze_commit': released['freeze_commit'],
              'repository_commit': receipt['repository_commit'],
              'diagnostic_sha256': hashlib.sha256(args.diagnostic.read_bytes()).hexdigest(),
              'workbook_decoders': ['openpyxl', 'raw OOXML via Python ElementTree'],
              'matching_standard_cells_checked': n_standards,
              'matching_response_cells_checked': n_responses,
              'matching_absorbance_cells_checked': n_standards+n_responses,
              'time_cells_checked': nt_standards+nt_responses,
              'time_decoder_tolerance_seconds': .00051,
              'maximum_time_decoder_difference_seconds': max(dt_standards, dt_responses),
              'calibration_method': 'centered covariance divided by centered sum of squares',
              'production_analysis_imported': False,
              'relative_tolerance': 1e-11, 'absolute_tolerance': 1e-10, **stats}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, allow_nan=False)+'\n')
    print(json.dumps(result))


if __name__ == '__main__':
    main()
